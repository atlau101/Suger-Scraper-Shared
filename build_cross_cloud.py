#!/usr/bin/env python3
"""Build one exact-key Azure + GCP prospecting deliverable.

The inputs remain authoritative and untouched. GCP's workbook-local joins use
its stored vendor_key first; cross-cloud identity is then recomputed with the
current shared canonicaliser so older published keys cannot hide overlaps.
"""

from __future__ import annotations

import argparse
import csv
import os
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from normalize import normalize_domain, vendor_key


UNIFIED_COLUMNS = [
    "vendor_key",
    "vendor_name",
    "website_domain",
    "marketplaces",
    "marketplace_count",
    "prospect_status",
    "review_reasons",
    "account_id",
    "account_name",
    "sdr_responsible",
    "sf_type",
    "sf_website",
    "source_row_count",
]

AUDIT_COLUMNS = [
    "source",
    "source_sheet",
    "source_row",
    "original_vendor_key",
    "current_vendor_key",
    "vendor_name",
    "key_changed",
    "merged_source_row_count",
]

BLOCKED_STATUSES = {
    "active_opportunity",
    "competitor",
    "customer",
    "existing_owned",
    "other_excluded",
}
REVIEW_STATUSES = {"existing_unknown", "match_review"}
ACTIONABLE_STATUSES = {"existing_unowned", "net_new", "net_new_candidate"}
KNOWN_STATUSES = BLOCKED_STATUSES | REVIEW_STATUSES | ACTIONABLE_STATUSES
BLOCKED_PRIORITY = [
    "active_opportunity",
    "customer",
    "competitor",
    "existing_owned",
    "other_excluded",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="D9EAF7")


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"CSV has no header: {path}")
        return list(reader.fieldnames), [dict(row) for row in reader]


def read_sheet_rows(workbook, name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if name not in workbook.sheetnames:
        raise ValueError(f"Required GCP sheet missing: {name}")
    values = workbook[name].iter_rows(values_only=True)
    raw_header = next(values, None)
    if not raw_header:
        raise ValueError(f"GCP sheet has no header: {name}")
    header = [text(value) for value in raw_header]
    rows = []
    for source_row, values_row in enumerate(values, start=2):
        if not any(value is not None for value in values_row):
            continue
        row = dict(zip(header, values_row))
        row["__source_row"] = source_row
        rows.append(row)
    return header, rows


def index_rows(rows: Iterable[dict[str, Any]], field: str) -> dict[str, list[dict[str, Any]]]:
    indexed: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = text(row.get(field))
        if key:
            indexed[key].append(row)
    return indexed


def attach_gcp_sheet(
    base_rows: list[dict[str, Any]],
    extra_rows: list[dict[str, Any]],
    prefix: str,
    extra_header: list[str],
) -> None:
    """Attach workbook-local data using the source workbook's stored key."""
    indexed = index_rows(extra_rows, "vendor_key")
    for row in base_rows:
        matches = indexed.get(text(row.get("vendor_key")), [])
        for column in extra_header:
            values = unique_values(matches, column)
            row[f"__{prefix}{column}"] = " | ".join(values)


def unique_values(rows: Iterable[dict[str, Any]], field: str) -> list[str]:
    seen: set[str] = set()
    values: list[str] = []
    for row in rows:
        value = text(row.get(field))
        if value and value not in seen:
            seen.add(value)
            values.append(value)
    return values


def combined_value(rows: Iterable[dict[str, Any]], field: str) -> str:
    return " | ".join(unique_values(rows, field))


def source_values(
    azure_rows: list[dict[str, Any]],
    gcp_rows: list[dict[str, Any]],
    field: str,
) -> list[str]:
    return unique_values([*azure_rows, *gcp_rows], field)


def resolve_status(statuses: set[str]) -> tuple[str, list[str]]:
    reasons: list[str] = []
    unknown = sorted(status for status in statuses if status not in KNOWN_STATUSES)
    if unknown:
        reasons.append("unknown_account_status:" + ",".join(unknown))

    blocked = statuses & BLOCKED_STATUSES
    review = statuses & REVIEW_STATUSES
    actionable = statuses & ACTIONABLE_STATUSES

    if review:
        reasons.append("unresolved_account_match")
        return "review", reasons
    if blocked and actionable:
        reasons.append("account_status_conflict")
        return "review", reasons
    if unknown:
        return "review", reasons
    if blocked:
        return next(status for status in BLOCKED_PRIORITY if status in blocked), reasons
    if "existing_unowned" in actionable:
        return "existing_unowned", reasons
    if actionable:
        return "net_new", reasons
    return "review", ["missing_account_status"]


def reconcile_field(
    azure_rows: list[dict[str, Any]],
    gcp_rows: list[dict[str, Any]],
    field: str,
) -> tuple[str, bool]:
    values = source_values(azure_rows, gcp_rows, field)
    normalized = {value.casefold() for value in values}
    if len(normalized) > 1:
        return "", True
    return (values[0] if values else ""), False


def reconcile_domain(
    azure_rows: list[dict[str, Any]], gcp_rows: list[dict[str, Any]]
) -> tuple[str, bool]:
    values = source_values(azure_rows, gcp_rows, "website_domain")
    normalized = {normalize_domain(value) for value in values if normalize_domain(value)}
    if len(normalized) > 1:
        return "", True
    return (next(iter(normalized)) if normalized else ""), False


def merge_prefixed(
    output: dict[str, Any],
    rows: list[dict[str, Any]],
    header: list[str],
    prefix: str,
) -> None:
    for column in header:
        output[f"{prefix}{column}"] = combined_value(rows, column)


def build_combined_rows(
    azure_rows: list[dict[str, Any]],
    gcp_rows: list[dict[str, Any]],
    azure_header: list[str],
    gcp_header: list[str],
    queue_header: list[str],
    audit_header: list[str],
) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]]]:
    groups: dict[str, dict[str, list[dict[str, Any]]]] = defaultdict(
        lambda: {"azure": [], "gcp": []}
    )
    source_audit: list[dict[str, Any]] = []

    for source, rows in (("Azure", azure_rows), ("GCP", gcp_rows)):
        for position, row in enumerate(rows, start=2):
            name = text(row.get("vendor_name"))
            current_key = vendor_key(name)
            if not current_key:
                raise ValueError(f"{source} row {position} has no usable vendor_name")
            row["__current_vendor_key"] = current_key
            row.setdefault("__source_row", position)
            groups[current_key][source.casefold()].append(row)

    data_columns = (
        UNIFIED_COLUMNS
        + [f"azure_{column}" for column in azure_header]
        + [f"gcp_{column}" for column in gcp_header]
        + [f"gcp_queue_{column}" for column in queue_header]
        + [f"gcp_audit_{column}" for column in audit_header]
    )

    combined: list[dict[str, Any]] = []
    for current_key, source_group in groups.items():
        azure_group = source_group["azure"]
        gcp_group = source_group["gcp"]
        all_rows = [*azure_group, *gcp_group]
        review_reasons: list[str] = []

        if len(azure_group) > 1:
            review_reasons.append("azure_identity_collision")
        if len(gcp_group) > 1:
            review_reasons.append("gcp_identity_collision")
        identity_collision = len(azure_group) > 1 or len(gcp_group) > 1

        statuses = {
            text(row.get("account_status"))
            for row in all_rows
            if text(row.get("account_status"))
        }
        prospect_status, status_reasons = resolve_status(statuses)
        review_reasons.extend(status_reasons)

        unified_sf: dict[str, str] = {}
        fatal_conflict = False
        for field in ("account_id", "account_name", "sdr_responsible", "sf_type", "sf_website"):
            value, conflict = reconcile_field(azure_group, gcp_group, field)
            unified_sf[field] = value
            if conflict:
                review_reasons.append(f"{field}_conflict")
                fatal_conflict = True

        website_domain, domain_conflict = reconcile_domain(azure_group, gcp_group)
        if domain_conflict:
            review_reasons.append("website_domain_conflict")
            fatal_conflict = True

        if fatal_conflict or identity_collision:
            prospect_status = "review"

        partner_review = any(
            text(row.get("partner_status")) == "match_review"
            or text(row.get("partner_lookup_status")) == "match_review"
            for row in all_rows
        )
        if partner_review:
            review_reasons.append("partner_identity_review")

        marketplaces = [name for name, rows in (("Azure", azure_group), ("GCP", gcp_group)) if rows]
        display_name = (
            unique_values(azure_group, "vendor_name")
            or unique_values(gcp_group, "vendor_name")
        )[0]
        output: dict[str, Any] = {
            "vendor_key": current_key,
            "vendor_name": display_name,
            "website_domain": website_domain,
            "marketplaces": "; ".join(marketplaces),
            "marketplace_count": len(marketplaces),
            "prospect_status": prospect_status,
            "review_reasons": "; ".join(dict.fromkeys(review_reasons)),
            **unified_sf,
            "source_row_count": len(all_rows),
        }
        merge_prefixed(output, azure_group, azure_header, "azure_")
        merge_prefixed(output, gcp_group, gcp_header, "gcp_")
        for column in queue_header:
            output[f"gcp_queue_{column}"] = " | ".join(
                value
                for value in dict.fromkeys(
                    text(row.get(f"__queue_{column}")) for row in gcp_group
                )
                if value
            )
        for column in audit_header:
            output[f"gcp_audit_{column}"] = " | ".join(
                value
                for value in dict.fromkeys(
                    text(row.get(f"__audit_{column}")) for row in gcp_group
                )
                if value
            )
        combined.append(output)

        for source, rows in (("Azure", azure_group), ("GCP", gcp_group)):
            for row in rows:
                old_key = text(row.get("vendor_key"))
                source_audit.append(
                    {
                        "source": source,
                        "source_sheet": "azure_prospects.csv" if source == "Azure" else "All Accounts",
                        "source_row": row["__source_row"],
                        "original_vendor_key": old_key,
                        "current_vendor_key": current_key,
                        "vendor_name": text(row.get("vendor_name")),
                        "key_changed": "true" if old_key != current_key else "false",
                        "merged_source_row_count": len(rows),
                    }
                )

    combined.sort(key=combined_sort_key)
    source_audit.sort(key=lambda row: (row["source"], int(row["source_row"])))
    return data_columns, combined, source_audit


def combined_sort_key(row: dict[str, Any]) -> tuple[Any, ...]:
    status_rank = {"existing_unowned": 0, "net_new": 1}.get(text(row.get("prospect_status")), 2)
    return (-int(row["marketplace_count"]), status_rank, text(row.get("vendor_name")).casefold())


def write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def add_table_sheet(
    workbook: Workbook,
    title: str,
    note: str,
    columns: list[str],
    rows: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append([note])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    sheet["A1"].fill = NOTE_FILL
    sheet["A1"].font = Font(bold=True, color="1F1F1F")
    sheet["A1"].alignment = Alignment(wrap_text=True)
    sheet.append(columns)
    for cell in sheet[2]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{max(2, sheet.max_row)}"
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 32
    for index, column in enumerate(columns, start=1):
        width = max(12, min(42, len(column) + 2))
        if column in {"vendor_name", "account_name", "review_reasons"} or column.endswith(("why", "sources")):
            width = 36
        sheet.column_dimensions[get_column_letter(index)].width = width


def add_readme(workbook: Workbook, counts: dict[str, int], azure_path: Path, gcp_path: Path) -> None:
    sheet = workbook.create_sheet("Read Me First")
    rows = [
        ("Item", "Detail"),
        ("Purpose", "Exact-key union of the supplied Azure and GCP prospecting snapshots."),
        ("All Accounts", str(counts["all"])),
        ("Prospecting Queue", str(counts["queue"])),
        ("Cross-Cloud Prospects", str(counts["cross_cloud"])),
        ("Match Review", str(counts["review"])),
        ("Identity", "vendor_key is recomputed with shared.normalize; no fuzzy name fallback."),
        ("Priority", "Cross-cloud first, then existing-unowned, then net-new; no synthetic ICP score."),
        ("Eligibility", "Only resolved net-new and existing-unowned rows enter Prospecting Queue."),
        ("Azure source", str(azure_path)),
        ("GCP source", str(gcp_path)),
        ("Raw listing evidence", "Remains in the source GCP workbook and is not copied here."),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 110
    for row in sheet.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def data_dictionary(columns: list[str]) -> list[dict[str, str]]:
    unified = {
        "vendor_key": "Current exact cross-cloud company key.",
        "vendor_name": "Display name; Azure value first when available.",
        "website_domain": "Normalized domain; blank when source domains conflict.",
        "marketplaces": "Marketplaces represented by the supplied scrape rows.",
        "marketplace_count": "Count derived from scrape presence, not Salesforce flags.",
        "prospect_status": "Conservative unified Salesforce eligibility status.",
        "review_reasons": "Semicolon-separated reasons requiring human review.",
        "account_id": "Unified only when nonblank source values agree.",
        "account_name": "Unified only when nonblank source values agree.",
        "sdr_responsible": "Unified only when nonblank source values agree.",
        "sf_type": "Unified only when nonblank source values agree.",
        "sf_website": "Unified only when nonblank source values agree.",
        "source_row_count": "Azure and GCP account rows merged into this company.",
    }
    rows = []
    for column in columns:
        if column in unified:
            source = "derived"
            meaning = unified[column]
        elif column.startswith("azure_"):
            source = "Azure azure_prospects.csv"
            meaning = f"Source field `{column[6:]}` preserved from Azure."
        elif column.startswith("gcp_queue_"):
            source = "GCP Prospecting Queue"
            meaning = f"Source field `{column[10:]}` attached by stored GCP vendor_key."
        elif column.startswith("gcp_audit_"):
            source = "GCP Evidence Audit"
            meaning = f"Source field `{column[10:]}` attached by stored GCP vendor_key."
        else:
            source = "GCP All Accounts"
            meaning = f"Source field `{column[4:]}` preserved from GCP."
        rows.append({"column": column, "meaning": meaning, "source": source})
    return rows


def write_workbook(
    path: Path,
    columns: list[str],
    rows: list[dict[str, Any]],
    source_audit: list[dict[str, Any]],
    azure_path: Path,
    gcp_path: Path,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    queue = [row for row in rows if row["prospect_status"] in {"net_new", "existing_unowned"}]
    cross_cloud = [row for row in rows if row["marketplace_count"] == 2]
    existing_unowned = [row for row in rows if row["prospect_status"] == "existing_unowned"]
    review = [row for row in rows if row["review_reasons"]]
    add_readme(
        workbook,
        {"all": len(rows), "queue": len(queue), "cross_cloud": len(cross_cloud), "review": len(review)},
        azure_path,
        gcp_path,
    )
    add_table_sheet(workbook, "Prospecting Queue", "Resolved net-new and existing-unowned companies.", columns, queue)
    add_table_sheet(workbook, "Cross-Cloud Prospects", "Companies present in both supplied scrapes; check prospect_status before outreach.", columns, cross_cloud)
    add_table_sheet(workbook, "Existing Unowned", "Resolved Salesforce accounts without an SDR owner.", columns, existing_unowned)
    add_table_sheet(workbook, "All Accounts", "Complete one-row-per-company union.", columns, rows)
    add_table_sheet(workbook, "Match Review", "Rows with identity, source, or matching issues to inspect.", columns, review)
    add_table_sheet(workbook, "Source Audit", "Every Azure and GCP account-level source row mapped to its current key.", AUDIT_COLUMNS, source_audit)
    add_table_sheet(
        workbook,
        "Data Dictionary",
        "Unified and source-prefixed column definitions.",
        ["column", "meaning", "source"],
        data_dictionary(columns),
    )
    workbook.save(path)


def build_artifacts(azure_path: Path, gcp_path: Path, output_prefix: Path) -> tuple[Path, Path]:
    xlsx_path = output_prefix.with_suffix(".xlsx")
    csv_path = output_prefix.with_suffix(".csv")
    existing = [path for path in (xlsx_path, csv_path) if path.exists()]
    if existing:
        raise FileExistsError("Refusing to overwrite: " + ", ".join(str(path) for path in existing))
    if not azure_path.is_file():
        raise FileNotFoundError(azure_path)
    if not gcp_path.is_file():
        raise FileNotFoundError(gcp_path)

    azure_header, azure_rows = read_csv_rows(azure_path)
    workbook = load_workbook(gcp_path, read_only=True, data_only=True)
    try:
        gcp_header, gcp_rows = read_sheet_rows(workbook, "All Accounts")
        queue_header, queue_rows = read_sheet_rows(workbook, "Prospecting Queue")
        audit_header, audit_rows = read_sheet_rows(workbook, "Evidence Audit")
    finally:
        workbook.close()

    attach_gcp_sheet(gcp_rows, queue_rows, "queue_", queue_header)
    attach_gcp_sheet(gcp_rows, audit_rows, "audit_", audit_header)
    columns, rows, source_audit = build_combined_rows(
        azure_rows, gcp_rows, azure_header, gcp_header, queue_header, audit_header
    )

    output_prefix.parent.mkdir(parents=True, exist_ok=True)
    tmp_csv = csv_path.with_name(csv_path.name + ".tmp")
    tmp_xlsx = xlsx_path.with_name(xlsx_path.name + ".tmp")
    try:
        write_csv(tmp_csv, columns, rows)
        write_workbook(tmp_xlsx, columns, rows, source_audit, azure_path, gcp_path)
        os.replace(tmp_csv, csv_path)
        os.replace(tmp_xlsx, xlsx_path)
    finally:
        for temporary in (tmp_csv, tmp_xlsx):
            if temporary.exists():
                temporary.unlink()
    return xlsx_path, csv_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--azure-csv", required=True, type=Path)
    parser.add_argument("--gcp-workbook", required=True, type=Path)
    parser.add_argument("--output-prefix", required=True, type=Path)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    xlsx_path, csv_path = build_artifacts(args.azure_csv, args.gcp_workbook, args.output_prefix)
    print(f"Wrote {xlsx_path}")
    print(f"Wrote {csv_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
