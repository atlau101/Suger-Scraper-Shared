import csv
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

SHARED = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SHARED))

from build_cross_cloud import (  # noqa: E402
    build_artifacts,
    build_combined_rows,
    resolve_status,
)
from normalize import vendor_key  # noqa: E402


AZURE_HEADER = [
    "vendor_key",
    "vendor_name",
    "website_domain",
    "account_status",
    "account_id",
    "account_name",
    "sdr_responsible",
    "sf_type",
    "sf_website",
    "signal_score",
    "verdict",
]
GCP_HEADER = [
    "vendor_key",
    "vendor_name",
    "website_domain",
    "account_status",
    "account_id",
    "account_name",
    "sdr_responsible",
    "sf_type",
    "sf_website",
    "partner_lookup_status",
]
QUEUE_HEADER = ["vendor_key", "vendor_name", "tier", "why"]
EVIDENCE_HEADER = ["vendor_key", "vendor_name", "key_evidence", "research_sources"]


def row(name, status, **values):
    base = {
        "vendor_key": values.pop("stored_key", vendor_key(name)),
        "vendor_name": name,
        "website_domain": values.pop("website_domain", ""),
        "account_status": status,
        "account_id": "",
        "account_name": "",
        "sdr_responsible": "",
        "sf_type": "",
        "sf_website": "",
    }
    base.update(values)
    return base


class ResolveStatusTests(unittest.TestCase):
    def test_conservative_status_resolution(self):
        self.assertEqual(resolve_status({"net_new"})[0], "net_new")
        self.assertEqual(resolve_status({"net_new_candidate", "existing_unowned"})[0], "existing_unowned")
        self.assertEqual(resolve_status({"customer"})[0], "customer")
        self.assertEqual(resolve_status({"existing_unknown"})[0], "review")
        status, reasons = resolve_status({"net_new", "existing_owned"})
        self.assertEqual(status, "review")
        self.assertIn("account_status_conflict", reasons)


class CombineRowsTests(unittest.TestCase):
    def test_exact_current_key_join_conflicts_and_collision(self):
        azure = [
            row(
                "Trafficguard Pty Ltd",
                "net_new",
                website_domain="trafficguard.ai",
                account_id="AZ-1",
                signal_score="80",
                verdict="strong",
            ),
            row("Acme", "net_new", website_domain="acme.com"),
            row("Owned Only", "existing_owned"),
        ]
        gcp = [
            row(
                "Trafficguard",
                "net_new_candidate",
                stored_key="legacy-key",
                website_domain="trafficguard.ai",
                account_id="GCP-2",
            ),
            row("Acmee", "net_new_candidate", website_domain="acmee.com"),
            row("APPLIVERY S.L.", "net_new_candidate", stored_key="old-applivery"),
            row("Applivery", "net_new_candidate"),
        ]
        columns, combined, audit = build_combined_rows(
            azure, gcp, AZURE_HEADER, GCP_HEADER, QUEUE_HEADER, EVIDENCE_HEADER
        )
        by_name = {item["vendor_name"]: item for item in combined}

        trafficguard = by_name["Trafficguard Pty Ltd"]
        self.assertEqual(trafficguard["marketplace_count"], 2)
        self.assertEqual(trafficguard["account_id"], "")
        self.assertEqual(trafficguard["prospect_status"], "review")
        self.assertIn("account_id_conflict", trafficguard["review_reasons"])

        self.assertIn("Acme", by_name)
        self.assertIn("Acmee", by_name)
        self.assertEqual(by_name["Acme"]["marketplace_count"], 1)
        self.assertEqual(by_name["Acmee"]["marketplace_count"], 1)

        applivery = next(item for item in combined if item["vendor_key"] == vendor_key("Applivery"))
        self.assertEqual(applivery["source_row_count"], 2)
        self.assertEqual(applivery["prospect_status"], "review")
        self.assertIn("gcp_identity_collision", applivery["review_reasons"])

        self.assertIn("azure_signal_score", columns)
        self.assertIn("gcp_partner_lookup_status", columns)
        self.assertEqual(len(audit), len(azure) + len(gcp))
        legacy = next(item for item in audit if item["original_vendor_key"] == "legacy-key")
        self.assertEqual(legacy["key_changed"], "true")


class ArtifactTests(unittest.TestCase):
    def test_writes_expected_workbook_and_preserves_inputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            azure_path = root / "azure.csv"
            gcp_path = root / "gcp.xlsx"
            prefix = root / "out" / "combined_v1"

            azure_rows = [
                row("Both Co", "net_new", website_domain="both.example", signal_score="70", verdict="possible"),
                row("Azure Co", "existing_unowned", website_domain="azure.example"),
            ]
            with azure_path.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.DictWriter(handle, fieldnames=AZURE_HEADER, extrasaction="ignore")
                writer.writeheader()
                writer.writerows(azure_rows)

            workbook = Workbook()
            all_accounts = workbook.active
            all_accounts.title = "All Accounts"
            all_accounts.append(GCP_HEADER)
            gcp_rows = [
                row("Both Company", "net_new_candidate", stored_key="stale", website_domain="both.example"),
                row("GCP Co", "customer", website_domain="gcp.example"),
            ]
            # Use an exact canonical overlap in the fixture.
            gcp_rows[0]["vendor_name"] = "Both Co"
            for item in gcp_rows:
                all_accounts.append([item.get(column, "") for column in GCP_HEADER])
            queue = workbook.create_sheet("Prospecting Queue")
            queue.append(QUEUE_HEADER)
            queue.append(["stale", "Both Co", "1", "Clear fit"])
            evidence = workbook.create_sheet("Evidence Audit")
            evidence.append(EVIDENCE_HEADER)
            evidence.append(["stale", "Both Co", "Paid listing", "https://example.test"])
            workbook.save(gcp_path)

            before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (azure_path, gcp_path)}
            xlsx_path, csv_path = build_artifacts(azure_path, gcp_path, prefix)
            after = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (azure_path, gcp_path)}
            self.assertEqual(before, after)

            with csv_path.open(newline="", encoding="utf-8-sig") as handle:
                csv_rows = list(csv.DictReader(handle))
            self.assertEqual(len(csv_rows), 3)
            both = next(item for item in csv_rows if item["vendor_name"] == "Both Co")
            self.assertEqual(both["marketplace_count"], "2")
            self.assertEqual(both["gcp_queue_tier"], "1")
            self.assertEqual(both["gcp_audit_key_evidence"], "Paid listing")

            output = load_workbook(xlsx_path, read_only=False, data_only=True)
            self.assertEqual(
                output.sheetnames,
                [
                    "Read Me First",
                    "Prospecting Queue",
                    "Cross-Cloud Prospects",
                    "Existing Unowned",
                    "All Accounts",
                    "Match Review",
                    "Source Audit",
                    "Data Dictionary",
                ],
            )
            self.assertEqual(output["All Accounts"].max_row - 2, len(csv_rows))
            for name in output.sheetnames[1:]:
                self.assertEqual(output[name].freeze_panes, "A3")
                self.assertTrue(output[name].auto_filter.ref)
            output.close()

            with self.assertRaises(FileExistsError):
                build_artifacts(azure_path, gcp_path, prefix)


if __name__ == "__main__":
    unittest.main()
